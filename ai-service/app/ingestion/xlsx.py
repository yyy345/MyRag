"""First-sheet XLSX tables, preserving the legacy raw-value contract."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import math
import posixpath
from xml.etree.ElementTree import ParseError as XMLParseError, fromstring
from xml.parsers.expat import ExpatError, ParserCreate
from zipfile import BadZipFile, LargeZipFile, ZipFile
import zlib

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.utils.datetime import time_to_days, timedelta_to_days, to_excel
from openpyxl.utils.exceptions import InvalidFileException

from app.ingestion.csv import markdown_cell
from app.ingestion.parser import JS_WHITESPACE, ParseError

MAX_XLSX_PARTS = 2000
MAX_XLSX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_XLSX_XML_BYTES = 16 * 1024 * 1024
MAX_XLSX_ROWS = 100_000
MAX_XLSX_COLUMNS = 1000
MAX_XLSX_CELLS = 500_000
MAX_XLSX_CHARACTERS = 2_000_000

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WORKBOOK_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


def invalid_xlsx():
    return ParseError("invalid_xlsx", "XLSX 文件无效或已损坏，请确认不是加密文件或旧版 XLS")


def reject_doctype(*_args):
    raise invalid_xlsx()


def number_text(value):
    # SheetJS exposes raw values as JavaScript Numbers, not formatted currency,
    # percentages or date strings. Normalize exponent spelling and thresholds.
    value = float(value)
    if not math.isfinite(value):
        raise invalid_xlsx()
    if value == 0:
        return "0"
    text = repr(value)
    if 1e-6 <= abs(value) < 1e21:
        text = format(Decimal(text), "f")
        return text.rstrip("0").rstrip(".") if "." in text else text
    mantissa, exponent = text.split("e")
    mantissa = mantissa.removesuffix(".0")
    return f"{mantissa}e{int(exponent):+d}"


def first_sheet_data(package):
    manifest = fromstring(package.read("[Content_Types].xml"))
    workbook_path = next((part.get("PartName", "").lstrip("/") for part in manifest if part.get("ContentType") == WORKBOOK_TYPE), None)
    if not workbook_path:
        raise invalid_xlsx()
    workbook = fromstring(package.read(workbook_path))
    sheet = workbook.find(f"{S}sheets/{S}sheet")
    if sheet is None:
        raise ParseError("empty_document", "XLSX 文件无工作表")
    parent, name = posixpath.split(workbook_path)
    relationships = fromstring(package.read(posixpath.join(parent, "_rels", name + ".rels")))
    relation = next((item for item in relationships if item.get("Id") == sheet.get(R + "id")), None)
    if relation is None or relation.get("TargetMode") == "External" or not relation.get("Type", "").endswith("/worksheet"):
        raise invalid_xlsx()
    target = relation.get("Target", "")
    sheet_path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(parent, target))
    if sheet_path.startswith("../") or sheet_path == "..":
        raise invalid_xlsx()
    return fromstring(package.read(sheet_path))


def sheet_bounds_and_numbers(sheet):
    """Validate sparse extents before openpyxl expands rows; retain date serials."""
    bounds = None
    numbers = {}
    dimension = sheet.find(S + "dimension")
    if dimension is not None and dimension.get("ref"):
        min_col, min_row, max_col, max_row = range_boundaries(dimension.get("ref"))
        if None in (min_col, min_row, max_col, max_row):
            raise invalid_xlsx()
        bounds = [min_row, min_col, max_row, max_col]
    previous_row = 0
    seen = set()
    for row in sheet.findall(f"{S}sheetData/{S}row"):
        row_number = int(row.get("r", previous_row + 1))
        if row_number <= previous_row:
            raise invalid_xlsx()
        previous_row = row_number
        previous_column = 0
        for cell in row.findall(S + "c"):
            position = coordinate_to_tuple(cell.get("r")) if cell.get("r") else (row_number, previous_column + 1)
            if position[0] != row_number or position[1] <= previous_column or position in seen:
                raise invalid_xlsx()
            previous_column = position[1]
            seen.add(position)
            if len(seen) > MAX_XLSX_CELLS:
                raise ParseError("xlsx_size_limit", "XLSX 单元格数量超过限制", 413)
            if bounds is None:
                bounds = [*position, *position]
            else:
                bounds = [min(bounds[0], position[0]), min(bounds[1], position[1]), max(bounds[2], position[0]), max(bounds[3], position[1])]
            value = cell.find(S + "v")
            if cell.get("t", "n") == "n" and value is not None and value.text:
                numbers[position] = number_text(value.text)
    if bounds is None:
        raise ParseError("empty_document", "XLSX 第一张工作表无有效内容")
    min_row, min_col, max_row, max_col = bounds
    if min_row < 1 or min_col < 1 or max_row < min_row or max_col < min_col or max_row > 1_048_576 or max_col > 16_384:
        raise invalid_xlsx()
    rows, columns = max_row - min_row + 1, max_col - min_col + 1
    if rows > MAX_XLSX_ROWS or columns > MAX_XLSX_COLUMNS or rows * columns > MAX_XLSX_CELLS:
        raise ParseError("xlsx_size_limit", "XLSX 行数、列数或单元格范围超过限制", 413)
    return bounds, numbers


def cell_text(cell, raw_numbers, epoch):
    value = cell.value
    if value is None:
        return ""
    position = (cell.row, cell.column)
    if position in raw_numbers:
        # Avoid date conversion rounding and Excel's serial-60 leap-year quirk.
        return raw_numbers[position]
    if cell.data_type == "e":
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return number_text(to_excel(value, epoch))
    if isinstance(value, time):
        return number_text(time_to_days(value))
    if isinstance(value, timedelta):
        return number_text(timedelta_to_days(value))
    if isinstance(value, (int, float)):
        return number_text(value)
    return str(value).replace("\x00", "")


def parse_xlsx(content: bytes) -> str:
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as package:
            parts = package.infolist()
            names = [part.filename for part in parts]
            if len(names) != len(set(names)):
                raise invalid_xlsx()
            if len(parts) > MAX_XLSX_PARTS or sum(part.file_size for part in parts) > MAX_XLSX_EXPANDED_BYTES:
                raise ParseError("xlsx_size_limit", "XLSX 解压内容过大或文件项过多", 413)
            for part in parts:
                if part.flag_bits & 1:
                    raise invalid_xlsx()
                if part.filename.lower().endswith((".xml", ".rels")):
                    if part.file_size > MAX_XLSX_XML_BYTES:
                        raise ParseError("xlsx_size_limit", "XLSX XML 内容过大", 413)
                    parser = ParserCreate()
                    parser.StartDoctypeDeclHandler = reject_doctype
                    parser.Parse(package.read(part), True)
            bounds, raw_numbers = sheet_bounds_and_numbers(first_sheet_data(package))
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        min_row, min_col, max_row, max_col = bounds
        lines = []
        characters = 0
        for row in sheet.iter_rows(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col):
            cells = [cell_text(cell, raw_numbers, workbook.epoch) for cell in row]
            if not any(value.strip(JS_WHITESPACE) for value in cells):
                continue
            line = " | ".join(markdown_cell(value) for value in cells)
            addition = [line] if lines else [line, " | ".join(["---"] * len(cells))]
            characters += sum(len(item) + 1 for item in addition)
            if characters > MAX_XLSX_CHARACTERS:
                raise ParseError("xlsx_text_limit", "XLSX 解析文本过长", 413)
            lines.extend(addition)
        if not lines:
            raise ParseError("empty_document", "XLSX 第一张工作表无有效内容")
        return "\n".join(lines)
    except ParseError:
        raise
    except (BadZipFile, LargeZipFile, XMLParseError, ExpatError, InvalidFileException, zlib.error, OSError, ValueError, KeyError, TypeError, AttributeError, IndexError, OverflowError, RecursionError, NotImplementedError):
        raise invalid_xlsx() from None
    finally:
        if workbook is not None:
            workbook.close()
