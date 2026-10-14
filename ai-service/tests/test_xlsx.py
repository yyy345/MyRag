from io import BytesIO
import unittest
from unittest.mock import patch
from uuid import uuid4
import warnings
from zipfile import ZipFile

from fastapi.testclient import TestClient
from app.main import create_app
from app.ingestion.parser import ParseError, parse_document
from xlsx_fixtures import FIXES, rewrite, samples


class XLSXParserTests(unittest.TestCase):
    def test_first_sheet_tables_empty_cells_and_ranges(self):
        expected = {
            "basic.XLSX": "姓名 | 年龄 | 备注\n--- | --- | ---\n张三🙂 | 18 | \n李四 | 20 |  前后空格 ",
            "shared_strings.xlsx": "姓名 | 备注\n--- | ---\n中文🙂 | 粗体尾段",
            "multiple.xlsx": "第一表 | 值\n--- | ---\n中文 | 1",
            "blank_rows.xlsx": "a | b | c\n--- | --- | ---\n0 | false | \n1 |  | ",
            "shifted.xlsx": "列一 | 列二\n--- | ---\n数据 | ",
            "merged.xlsx": "总标题 | \n--- | ---\na | b",
            **{name: text for name, text in FIXES.items() if text is not None},
        }
        documents = samples()
        for name, text in expected.items():
            with self.subTest(name=name):
                self.assertEqual(parse_document(documents[name], name), text)

    def test_raw_dates_numbers_and_cached_formulas(self):
        expected = {
            "dates.xlsx": "日期 | 时间 | 闰日序列 | 精度 | 百分数\n--- | --- | --- | --- | ---\n45292.5 | 0.5208333333333334 | 60 | 45292.123456789 | 0.125",
            "date_1904.xlsx": "日期 | 日期序列\n--- | ---\n1 | 0",
            "out_of_range_date.xlsx": "原始数值\n---\n1e+22",
            "iso_dates.xlsx": "日期\n---\n45292.5",
            "numbers.xlsx": "值 | 布尔\n--- | ---\n0 | false\n0 | true\n1.5 | \n1e-7 | \n0.000001 | \n100000000000000000000 | \n1e+21 | \n9007199254740992 | ",
            "formula.xlsx": "输入 | 求和 | 文本 | 缺缓存 | 错误 | 布尔 | 日期\n--- | --- | --- | --- | --- | --- | ---\n1 | 3 | 中文 |  |  | true | 60",
        }
        documents = samples()
        for name, text in expected.items():
            with self.subTest(name=name), patch("socket.socket", side_effect=AssertionError("parser must not use network")):
                self.assertEqual(parse_document(documents[name], name), text)

    def test_empty_and_invalid_packages_fail_safely(self):
        documents = samples()
        for name, code in {
            "blank.xlsx": "empty_document", "empty_first.xlsx": "empty_document",
            "corrupt.xlsx": "invalid_xlsx", "fake.xlsx": "invalid_xlsx",
            "missing.xlsx": "invalid_xlsx", "malformed.xlsx": "invalid_xlsx",
        }.items():
            with self.subTest(name=name), self.assertRaises(ParseError) as caught:
                parse_document(documents[name], name)
            self.assertEqual(caught.exception.code, code)
            self.assertEqual(caught.exception.status, 422)
            self.assertNotIn("SECOND_SHEET_PRIVATE", str(caught.exception))

    def test_limits_reject_before_large_row_expansion(self):
        documents = samples()
        for limit in ("MAX_XLSX_PARTS", "MAX_XLSX_EXPANDED_BYTES", "MAX_XLSX_XML_BYTES", "MAX_XLSX_ROWS", "MAX_XLSX_COLUMNS", "MAX_XLSX_CELLS", "MAX_XLSX_CHARACTERS"):
            with self.subTest(limit=limit), patch("app.ingestion.xlsx." + limit, 1), self.assertRaises(ParseError) as caught:
                parse_document(documents["basic.XLSX"], "large.xlsx")
            self.assertEqual(caught.exception.status, 413)
        oversized = rewrite(documents["basic.XLSX"], "xl/worksheets/sheet1.xml", lambda xml: xml.replace(b'A1:C3', b'A1:XFD1048576'))
        with patch("app.ingestion.xlsx.load_workbook") as loader, self.assertRaises(ParseError) as caught:
            parse_document(oversized, "oversized.xlsx")
        self.assertEqual(caught.exception.status, 413)
        loader.assert_not_called()

    def test_duplicate_parts_dtd_and_external_sheet_are_rejected(self):
        basic = samples()["basic.XLSX"]
        with ZipFile(BytesIO(basic)) as package:
            parts = {part.filename: package.read(part) for part in package.infolist()}
        duplicate = BytesIO()
        with warnings.catch_warnings(), ZipFile(duplicate, "w") as package:
            warnings.simplefilter("ignore", UserWarning)
            for name, body in parts.items():
                package.writestr(name, body)
            package.writestr("xl/workbook.xml", parts["xl/workbook.xml"])
        encrypted = bytearray(basic)
        for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
            start = encrypted.index(signature)
            encrypted[start + flag_offset] |= 1
        hostile = [
            duplicate.getvalue(),
            bytes(encrypted),
            rewrite(basic, "xl/worksheets/sheet1.xml", lambda xml: b'<!DOCTYPE worksheet [<!ENTITY leak SYSTEM "file:///private">]>' + xml[xml.index(b"<worksheet"):]),
            rewrite(basic, "xl/_rels/workbook.xml.rels", lambda xml: xml.replace(b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"', b'TargetMode="External" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"')),
        ]
        for document in hostile:
            with self.assertRaises(ParseError) as caught:
                parse_document(document, "hostile.xlsx")
            self.assertEqual(caught.exception.code, "invalid_xlsx")
            self.assertEqual(caught.exception.status, 422)

    def test_authenticated_multipart_and_safe_errors(self):
        token = "local-xlsx-test-" + "x" * 32
        documents = samples()
        with TestClient(create_app(token)) as client:
            for name in ("basic.XLSX", "dates.xlsx", "empty_first.xlsx", "corrupt.xlsx"):
                request_id = str(uuid4())
                response = client.post("/internal/v1/parse", headers={"Authorization": "Bearer " + token}, data={"request_id": request_id}, files={"file": (name, documents[name], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
                self.assertEqual(response.status_code, 200 if name in ("basic.XLSX", "dates.xlsx") else 422)
                if response.status_code == 200:
                    self.assertEqual(response.json()["request_id"], request_id)
                    self.assertEqual(response.json()["metadata"]["filename"], name)
                    self.assertEqual(response.json()["metadata"]["characters"], len(response.json()["text"]))
                else:
                    self.assertIn(response.json()["detail"]["code"], ("empty_document", "invalid_xlsx"))
                self.assertNotIn(token, response.text)

    def test_workbook_closed_after_success_and_text_limit(self):
        from openpyxl import load_workbook
        basic = samples()["basic.XLSX"]
        for limit in (2_000_000, 1):
            workbooks = []

            def load(*args, **kwargs):
                workbook = load_workbook(*args, **kwargs)
                workbooks.append(workbook)
                return workbook

            with patch("app.ingestion.xlsx.load_workbook", side_effect=load), patch("app.ingestion.xlsx.MAX_XLSX_CHARACTERS", limit):
                if limit == 1:
                    with self.assertRaises(ParseError):
                        parse_document(basic, "limit.xlsx")
                else:
                    parse_document(basic, "normal.xlsx")
            self.assertEqual(len(workbooks), 1)
            self.assertIsNone(workbooks[0]._archive.fp)


if __name__ == "__main__":
    unittest.main()
