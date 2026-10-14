"""Generated XLSX files; no Excel installation or downloaded user documents."""
from datetime import datetime, time
from io import BytesIO
from xml.etree.ElementTree import Element, SubElement, fromstring, tostring
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def archive(parts):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return output.getvalue()


def rewrite(content, name, transform):
    with ZipFile(BytesIO(content)) as package:
        parts = {part.filename: package.read(part) for part in package.infolist()}
    parts[name] = transform(parts[name])
    return archive(parts)


def document(rows, configure=None):
    workbook = Workbook()
    workbook.active.title = "数据"
    for row in rows:
        workbook.active.append(row)
    if configure:
        configure(workbook)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def cached_formula(content):
    def inject(xml):
        root = fromstring(xml)
        for coordinate, value, kind in [("B2", "3", "n"), ("C2", "中文", "str"), ("G2", "60", "n")]:
            cell = root.find(f"{S}sheetData/{S}row/{S}c[@r='{coordinate}']")
            cell.set("t", kind)
            cell.find(S + "v").text = value
        return tostring(root, encoding="utf-8")
    return rewrite(content, "xl/worksheets/sheet1.xml", inject)


def shared_strings():
    content = document([["姓名", "备注"], ["中文🙂", "粗体尾段"]])
    with ZipFile(BytesIO(content)) as package:
        parts = {part.filename: package.read(part) for part in package.infolist()}
    root = fromstring(parts["xl/worksheets/sheet1.xml"])
    values = ["姓名", "备注", "中文🙂", "粗体尾段"]
    for cell, index in zip(root.findall(f"{S}sheetData/{S}row/{S}c"), range(4)):
        for child in list(cell):
            cell.remove(child)
        cell.set("t", "s")
        SubElement(cell, S + "v").text = str(index)
    strings = Element(S + "sst", {"count": "4", "uniqueCount": "4"})
    for index, value in enumerate(values):
        item = SubElement(strings, S + "si")
        if index == 3:
            for piece in ("粗体", "尾段"):
                SubElement(SubElement(item, S + "r"), S + "t").text = piece
        else:
            SubElement(item, S + "t").text = value
    types = fromstring(parts["[Content_Types].xml"])
    SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
        "PartName": "/xl/sharedStrings.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    })
    parts["xl/sharedStrings.xml"] = tostring(strings, encoding="utf-8")
    parts["xl/worksheets/sheet1.xml"] = tostring(root, encoding="utf-8")
    parts["[Content_Types].xml"] = tostring(types, encoding="utf-8")
    return archive(parts)


def samples():
    def multiple(workbook):
        workbook.create_sheet("不应导入").append(["SECOND_SHEET_PRIVATE"])
        workbook.active = 1

    def date_formats(workbook):
        workbook.active["C2"].number_format = "yyyy-mm-dd"
        workbook.active["D2"].number_format = "yyyy-mm-dd hh:mm:ss.000"
        workbook.active["E2"].number_format = "0.00%"

    def shifted(workbook):
        workbook.active.delete_rows(1, 2)
        workbook.active["C4"] = "列一"
        workbook.active["D4"] = "列二"
        workbook.active["C6"] = "数据"

    def merged(workbook):
        workbook.active.merge_cells("A1:B1")

    def formula_date(workbook):
        workbook.active["G2"].number_format = "yyyy-mm-dd"

    def wrong_dimension(xml):
        root = fromstring(xml)
        root.find(S + "dimension").set("ref", "A1:A1")
        return tostring(root, encoding="utf-8")

    basic = document([["姓名", "年龄", "备注"], ["张三🙂", 18, None], ["李四", 20, " 前后空格 "]])
    formula = cached_formula(document([
        ["输入", "求和", "文本", "缺缓存", "错误", "布尔", "日期"],
        [1, "=SUM(A2,2)", '="中文"', "=1+2", "#DIV/0!", True, "=DATE(1900,2,29)"],
    ], formula_date))
    return {
        "basic.XLSX": basic,
        "shared_strings.xlsx": shared_strings(),
        "multiple.xlsx": document([["第一表", "值"], ["中文", 1]], multiple),
        "empty_first.xlsx": document([], multiple),
        "blank.xlsx": document([[None, " "], [None, None]]),
        "dates.xlsx": document([
            ["日期", "时间", "闰日序列", "精度", "百分数"],
            [datetime(2024, 1, 1, 12), time(12, 30), 60, 45292.123456789, 0.125],
        ], date_formats),
        "date_1904.xlsx": document([["日期", "日期序列"], [datetime(1904, 1, 2), 0]], lambda wb: setattr(wb, "epoch", CALENDAR_MAC_1904)),
        "out_of_range_date.xlsx": document([["原始数值"], [1e22]], lambda wb: setattr(wb.active["A2"], "number_format", "yyyy-mm-dd")),
        "iso_dates.xlsx": document([["日期"], [datetime(2024, 1, 1, 12)]], lambda wb: setattr(wb, "iso_dates", True)),
        "numbers.xlsx": document([["值", "布尔"], [0, False], [-0.0, True], [1.5, None], [1e-7, None], [1e-6, None], [1e20, None], [1e21, None], [9007199254740993, None]]),
        "formula.xlsx": formula,
        "blank_rows.xlsx": document([[None, None], ["a", "b", "c"], [None, None], [0, False, None], [1], [None, "  "]]),
        "shifted.xlsx": document([], shifted),
        "merged.xlsx": document([["总标题", None], ["a", "b"]], merged),
        "markdown.xlsx": document([["名称", "备注"], ["中文", "a|b\\c\n第二行"]]),
        "wrong_dimension.xlsx": rewrite(basic, "xl/worksheets/sheet1.xml", wrong_dimension),
        "corrupt.xlsx": b"PK\x03\x04broken",
        "fake.xlsx": b"not a workbook",
        "missing.xlsx": archive({"readme.txt": "missing workbook"}),
        "malformed.xlsx": rewrite(basic, "xl/worksheets/sheet1.xml", lambda _: b"<broken"),
    }


# These are deliberate changes, checked separately from exact legacy parity.
FIXES = {
    "markdown.xlsx": "名称 | 备注\n--- | ---\n中文 | a\\|b\\\\c ↵ 第二行",
    "wrong_dimension.xlsx": "姓名 | 年龄 | 备注\n--- | --- | ---\n张三🙂 | 18 | \n李四 | 20 |  前后空格 ",
    "fake.xlsx": None,
    "iso_dates.xlsx": "日期\n---\n45292.5",
}
